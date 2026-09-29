#table_constructor.py

#====================================
# 1. Librerías
#====================================

import file_extractor_destajo as fe
import pandas as pd
import openpyxl
import numpy as np
from pathlib import Path
import re

#====================================
# 2. Funciones Privadas (Uso Interno)
#====================================

def _build_table(path):
    #Esta función construye un dataframe con los datos correspondientes del encabezado y los datos de la tabla 
    headers_data, table_data = fe.extracted_data_destajo(path)

    for column, value in headers_data.items():
        table_data[column] = value

    return table_data

def _last_id(path):
    #Esta función obtiene el último id del archivo final de registro
    file_path = Path(path) / "Caratula de destajos.xlsx"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb["Registro"]
    column = "A"

    last_val = None
    # Recorremos desde la última fila hacia arriba para ignorar celdas vacías al final
    for row in range(sheet.max_row, 0, -1):
        cell_value = sheet[f"{column}{row}"].value
        if cell_value is not None:
            if isinstance(cell_value, str) and not re.search(r'\d+', cell_value):
                continue
            last_val = cell_value
            break

    wb.close()
    return last_val

def _assign_id_column(df, path):
    #Esta función asigna un id a cada fila del dataframe, en la columna ya existente

    start_id = 1
    last_id = _last_id(path)

    try:
        if last_id is not None:
            #Extrae la parte numérica
            numeric = re.search(r'\d+', str(last_id))
            if numeric: 
                start_id = int(numeric.group()) + 1 if numeric else 1
    except (ValueError, AttributeError):
        start_id = 1

    seq = np.arange(start_id, start_id + len(df))

    df["ID"] = "DEST_0" + seq.astype(str)

    return df

def _add_id(df):
    df["ID"] = None
    columnas = list(df.columns)
    columnas.insert(0, columnas.pop(columnas.index("ID")))
    return df[columnas]

def _fill_excel_file(df_table, path):
    #Esta función agrega información al archivo de excel de registro de destajos
    sheet_name = "Registro"
    file_path = Path(path)/"Caratula de destajos.xlsx"
    with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
        last_row = writer.sheets[sheet_name].max_row
        df_table.to_excel(writer, sheet_name=sheet_name, startrow=last_row, index=False, header=False)

#====================================
# 3. Función Pública (Interfaz)
#====================================

def process_excel_file(path, source_name):
    #Esta función procesa un archivo de excel, extrae la información relevante y la agrega a un archivo de registro
    file_path = Path(path)/source_name
    table_data = _build_table(file_path)
    table_data = _add_id(table_data)
    table_data = _assign_id_column(table_data, path)
    #table_data = _calculate_importe(table_data)
    _fill_excel_file(table_data, path)

__all__ = ["process_excel_file"]
