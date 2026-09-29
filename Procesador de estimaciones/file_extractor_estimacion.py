#file_extractor_estimacion.py
#Extractor para el "FORMATO ESTIMACION C.O." -> hoja Caratula (celdas amarillas + otros campos) + hoja Sabana Estimaciónes

#====================================
# 1. Librerías y Variables
#====================================

import re
import unicodedata
import openpyxl
import pandas as pd
from pathlib import Path

SHEET_CARATULA = "Caratula"
SHEET_SABANA = "Sabana Estimaciónes"

#Celdas de la Caratula a extraer. Las primeras 5 están resaltadas en amarillo
#(validado por color de relleno real, FFFFFF00); las últimas 4 no están
#resaltadas pero también se requieren.
interest_cells = {
    "NOMBRE_OBRA": "F9",            #amarillo
    "NO_ESTIMACION": "K9",          #amarillo
    "PERIODO_ESTIMACION": "K11",    #amarillo
    "CONTRATISTA": "K16",           #amarillo
    "TRABAJOS_A_REALIZAR": "K17",   #amarillo
    "NO_OBRA": "F10",
    "NO_CONTRATO": "F11",
    "FECHA_ELABORACION": "K10",
    "IMPORTE_CONTRATO_TOTAL": "K12",
}

#La tabla de conceptos en Sabana Estimaciónes puede tener hasta ~100 filas
#(varía por obra), por eso se lee dinámicamente desde TABLE_START hasta
#encontrar la palabra STOP_MARKER en la columna de CLAVE.
TABLE_START = 12
STOP_MARKER = "SUBTOTAL"
HEADER_ROW = 9   #fila con los títulos de cada recuadro (CATÁLOGO, AVANCE GENERAL, ESTIMACION N, POR EJERCER)
SUBHEADER_ROW = 11   #fila con los subtítulos de cada recuadro (Clave, Cantidad, %, Importe...)

#Nombres de columna para los campos base del catálogo de conceptos (recuadro
#izquierdo, siempre fijo).
#
#Los recuadros "ESTIMACION N" (1, 2, 3...) se IGNORAN a propósito: el
#recuadro "AVANCE GENERAL EJECUTADO" ya es acumulado (se validó que
#AVANCE_ACUM == ESTIMACION_N de la última columna, y que
#AVANCE_ACUM + POR_EJERCER == CANT_CONTRATO). Por eso cada archivo, sin
#importar cuántos recuadros ESTIMACION traiga adentro, siempre entrega
#exactamente el mismo esquema: catálogo + acumulado + por ejercer. El detalle
#de cuánto correspondió a cada periodo se puede recuperar después restando
#AVANCE_ACUM entre dos snapshots (archivos) consecutivos de la misma obra.
CATALOGO_MAP = {
    "Clave": "CLAVE",
    "Concepto": "CONCEPTO",
    "Unidad": "UNIDAD",
    "Cantidad": "CANT_CONTRATO",
    "P. U.": "P_UNITARIO",
    "Importe": "IMPORTE_CONTRATO",
}

#Subtítulos genéricos usados en los demás recuadros
SUBCOL_MAP = {
    "Cantidad": "CANTIDAD",
    "%": "PCT",
    "Importe": "IMPORTE",
}

#====================================
# 2. Funciones Privadas (Uso Interno)
#====================================

def _strip_accents(text):
    #Normaliza texto quitando acentos, para comparar etiquetas sin depender de tildes
    nfkd = unicodedata.normalize("NFKD", text)
    return "".join(c for c in nfkd if not unicodedata.combining(c)).upper().strip()


def _extract_header_data(sheet_caratula):
    #Extrae las celdas de interés (amarillas + adicionales) de la Caratula
    header_data = {}
    for column, cell in interest_cells.items():
        header_data[column] = sheet_caratula[cell].value
    return header_data


def _find_blocks(sheet):
    #Localiza cada recuadro de la tabla (celdas combinadas en HEADER_ROW) con su
    #etiqueta y rango de columnas. Esto permite detectar automáticamente
    #recuadros nuevos como "ESTIMACION 2", "ESTIMACION 3", etc.
    blocks = []
    for mc in sheet.merged_cells.ranges:
        if mc.min_row <= HEADER_ROW <= mc.max_row:
            label = sheet.cell(row=mc.min_row, column=mc.min_col).value
            if label:
                blocks.append({
                    "label": str(label).strip(),
                    "col_start": mc.min_col,
                    "col_end": mc.max_col,
                })
    blocks.sort(key=lambda b: b["col_start"])
    return blocks


def _block_subcolumns(sheet, block):
    #Lee los subtítulos (fila SUBHEADER_ROW) dentro del rango de columnas de un recuadro
    subcols = {}
    for col in range(block["col_start"], block["col_end"] + 1):
        value = sheet.cell(row=SUBHEADER_ROW, column=col).value
        if value:
            subcols[str(value).strip()] = col
    return subcols


def _build_column_plan(sheet):
    #Construye el plan de columnas a extraer: fijas para el catálogo de conceptos,
    #más los dos únicos recuadros de interés (AVANCE GENERAL y POR EJERCER).
    #Cualquier recuadro "ESTIMACION N" se detecta pero se descarta a propósito:
    #ver nota junto a CATALOGO_MAP.
    blocks = _find_blocks(sheet)

    catalogo_cols = None
    fixed_plan = []   #lista de (nombre_columna_final, columna_excel)
    found_avance = False
    found_por_ejercer = False

    for block in blocks:
        label_norm = _strip_accents(block["label"])
        subcols = _block_subcolumns(sheet, block)

        if "CATALOGO" in label_norm:
            catalogo_cols = {
                CATALOGO_MAP[sub]: col
                for sub, col in subcols.items()
                if sub in CATALOGO_MAP
            }
            continue

        if "ESTIMACION" in label_norm:
            #Recuadro de periodo individual: se ignora a propósito (ver nota arriba)
            continue

        if "AVANCE GENERAL" in label_norm:
            prefix = "AVANCE_ACUM"
            found_avance = True
        elif "POR EJERCER" in label_norm:
            prefix = "POR_EJERCER"
            found_por_ejercer = True
        else:
            #Recuadro no reconocido: se conserva por si el formato agrega
            #en el futuro otro recuadro fijo que no sea "ESTIMACION N"
            prefix = re.sub(r"[^A-Z0-9]+", "_", label_norm).strip("_")

        for sub, col in subcols.items():
            suffix = SUBCOL_MAP.get(sub, re.sub(r"[^A-Z0-9]+", "_", sub.upper()).strip("_"))
            fixed_plan.append((f"{prefix}_{suffix}", col))

    if catalogo_cols is None:
        raise ValueError("No se encontró el recuadro 'CATÁLOGO DE CONCEPTOS DE CONTRATO' en la hoja.")
    if not found_avance:
        raise ValueError("No se encontró el recuadro 'AVANCE GENERAL EJECUTADO' en la hoja.")
    if not found_por_ejercer:
        raise ValueError("No se encontró el recuadro 'POR EJERCER' en la hoja.")

    return catalogo_cols, fixed_plan


def _is_stop_row(sheet, row, clave_col):
    #La fila de fin de tabla trae STOP_MARKER en la columna de CLAVE
    value = sheet.cell(row=row, column=clave_col).value
    return isinstance(value, str) and STOP_MARKER in value.upper()


def _extract_table_data(sheet):
    #Recorre la tabla desde TABLE_START fila por fila, tomando solo las filas
    #donde existan CLAVE y CONCEPTO (evita notas sueltas como "100 conceptos"),
    #hasta toparse con la fila de SUBTOTAL. El plan de columnas (catálogo +
    #avance acumulado + por ejercer) es siempre el mismo tamaño sin importar
    #cuántos recuadros ESTIMACION_N tenga el archivo por dentro, porque esos
    #se descartan al construir el plan (_build_column_plan).
    catalogo_cols, fixed_plan = _build_column_plan(sheet)
    clave_col = catalogo_cols["CLAVE"]
    concepto_col = catalogo_cols["CONCEPTO"]

    rows = []
    row = TABLE_START

    while row <= sheet.max_row:
        if _is_stop_row(sheet, row, clave_col):
            break

        clave = sheet.cell(row=row, column=clave_col).value
        concepto = sheet.cell(row=row, column=concepto_col).value

        if clave is not None and concepto is not None:
            registro = {name: sheet.cell(row=row, column=col).value for name, col in catalogo_cols.items()}
            for name, col in fixed_plan:
                registro[name] = sheet.cell(row=row, column=col).value
            rows.append(registro)

        row += 1

    return pd.DataFrame(rows)

#====================================
# 3. Función Pública (Interfaz)
#====================================

def extracted_data_estimacion(path):
    #Función principal: abre el archivo, extrae encabezado (Caratula) y tabla (Sabana Estimaciónes)

    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_caratula = wb[SHEET_CARATULA]
    sheet_sabana = wb[SHEET_SABANA]

    header_data = _extract_header_data(sheet_caratula)
    table_data = _extract_table_data(sheet_sabana)

    wb.close()
    return header_data, table_data


__all__ = ["extracted_data_estimacion"]
